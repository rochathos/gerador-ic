from datetime import datetime
from pathlib import Path
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.core.config import settings
from app.core.logger import logger
from app.models.commit import Commit
from app.models.meeting import Meeting
from app.models.catalog_item import CatalogItem
from app.models.execution_history import ExecutionHistory


class ReportService:
    """Generates analytical Excel and PDF reports from execution data."""

    @classmethod
    def generate_excel(
        cls,
        history: ExecutionHistory,
        commits: List[Commit],
        meetings: Optional[List[Meeting]] = None,
        catalog_items: Optional[List[CatalogItem]] = None,
    ) -> Path:
        """Export execution data to an Excel (.xlsx) file."""
        wb = Workbook()
        # Tab 1: Resumo
        ws_summary = wb.active
        ws_summary.title = "Resumo da Execução"

        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=16, bold=True, color="1F2937")

        ws_summary["A1"] = "Productivity Assistant - Relatório de Execução"
        ws_summary["A1"].font = title_font

        ws_summary.append([])
        summary_rows = [
            ("ID da Execução", history.id),
            ("Período Inicial", history.start_date.strftime("%d/%m/%Y %H:%M")),
            ("Período Final", history.end_date.strftime("%d/%m/%Y %H:%M")),
            ("Repositório", history.repo_path or "N/D"),
            ("Status", history.status.upper()),
            ("Total de Commits", len(commits)),
            ("Total de Reuniões", len(meetings or [])),
            ("Total de ICs", len(catalog_items or [])),
            ("Data de Geração", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
        ]
        for label, val in summary_rows:
            ws_summary.append([label, val])

        for col in ["A", "B"]:
            ws_summary.column_dimensions[col].width = 30

        # Tab 2: Commits
        ws_commits = wb.create_sheet(title="Commits")
        headers = ["Hash", "Data/Hora", "Autor", "Mensagem", "Arquivos Alterados", "XMLs Alterados", "Linhas XML (+/-)", "Repositório"]
        ws_commits.append(headers)
        for col_num, _ in enumerate(headers, 1):
            cell = ws_commits.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

        for c in commits:
            if isinstance(c.files_changed, list):
                formatted_files = []
                for f in c.files_changed:
                    if isinstance(f, dict):
                        ins = f.get('insertions', 0)
                        dels = f.get('deletions', 0)
                        formatted_files.append(f"{f.get('path', '')} (+{ins}/-{dels})")
                    else:
                        formatted_files.append(str(f))
                files_str = ", ".join(formatted_files)
            else:
                files_str = str(c.files_changed or "")

            xml_info = f"{c.xml_files_count} arq" if c.has_xml_changes else "0"
            xml_diffs = f"+{c.xml_insertions} / -{c.xml_deletions}" if c.has_xml_changes else "0"

            ws_commits.append([
                c.short_hash,
                c.commit_date.strftime("%d/%m/%Y %H:%M"),
                c.author,
                c.message,
                files_str,
                xml_info,
                xml_diffs,
                c.repo_name or "",
            ])

        for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
            ws_commits.column_dimensions[col].width = 25

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = settings.REPORTS_DIR / f"relatorio_exec_{history.id}_{timestamp}.xlsx"
        settings.REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        wb.save(filepath)
        logger.info(f"Relatório Excel gerado: {filepath}")
        return filepath

    @classmethod
    def generate_pdf(
        cls,
        history: ExecutionHistory,
        commits: List[Commit],
        meetings: Optional[List[Meeting]] = None,
        catalog_items: Optional[List[CatalogItem]] = None,
    ) -> Path:
        """Export executive summary and commit details to a PDF file."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = settings.REPORTS_DIR / f"relatorio_exec_{history.id}_{timestamp}.pdf"
        settings.REPORTS_DIR.mkdir(parents=True, exist_ok=True)

        doc = SimpleDocTemplate(str(filepath), pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#1e293b"),
            spaceAfter=12,
        )
        heading_style = ParagraphStyle(
            "HeadingStyle",
            parent=styles["Heading2"],
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#3b82f6"),
            spaceBefore=14,
            spaceAfter=8,
        )
        normal_style = styles["Normal"]

        elements = [
            Paragraph("Productivity Assistant", title_style),
            Paragraph(f"<b>Relatório de Atividades - Execução #{history.id}</b>", styles["Heading3"]),
            Spacer(1, 10),
            Paragraph(
                f"<b>Período Analisado:</b> {history.start_date.strftime('%d/%m/%Y %H:%M')} até {history.end_date.strftime('%d/%m/%Y %H:%M')}<br/>"
                f"<b>Repositório:</b> {history.repo_path or 'Geral'}<br/>"
                f"<b>Total de Commits:</b> {len(commits)} | <b>Reuniões:</b> {len(meetings or [])} | <b>ICs:</b> {len(catalog_items or [])}<br/>"
                f"<b>Status:</b> {history.status.upper()}",
                normal_style,
            ),
            Spacer(1, 15),
            Paragraph("Commits Encontrados", heading_style),
        ]

        # Commits table
        table_data = [["Hash", "Data", "Autor", "Mensagem"]]
        for c in commits[:50]:  # Limit to 50 for layout readability
            short_msg = c.message.split("\n")[0][:60]
            table_data.append([
                c.short_hash,
                c.commit_date.strftime("%d/%m/%Y %H:%M"),
                c.author[:20],
                Paragraph(short_msg, normal_style),
            ])

        table = Table(table_data, colWidths=[60, 95, 110, 255])
        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ])
        )
        elements.append(table)

        doc.build(elements)
        logger.info(f"Relatório PDF gerado: {filepath}")
        return filepath
